"""Versiones PowerPoint editables del informe (español y portugués).

Uso:  .venv\\Scripts\\python.exe scripts\\21_informe_pptx.py              # genera los dos .pptx
      .venv\\Scripts\\python.exe scripts\\21_informe_pptx.py --coletar    # vuelca los textos a traducir

No redibuja ni reescribe el informe: ejecuta las funciones de página de 18_informe_pdf.py contra un
lienzo que REGISTRA cada elemento (cadenas, párrafos, tablas, imágenes, rectángulos y líneas, con sus
posiciones, fuentes y colores) y lo reproduce en PowerPoint con cuadros de texto y tablas nativos.
La versión en portugués usa los mismos elementos con los textos traducidos (scripts/traducao_pt.py)
y figuras regeneradas en portugués por 17_figuras.py en produtos/figuras_pt/.
Tipografía: IBM Plex Sans (assets/fonts/); si no está instalada, PowerPoint la sustituye.
"""
from __future__ import annotations

import importlib.util
import json
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from reportlab.pdfbase.pdfmetrics import stringWidth

from comum import caminho

RAIZ = caminho(".")
SALIDA = {"es": caminho("produtos/informe_sevilla_oferta_autobuses.pptx"),
          "pt": caminho("produtos/relatorio_sevilla_oferta_onibus_pt.pptx")}
FIG_PT = caminho("produtos/figuras_pt")
FUENTE = {"Plex": ("IBM Plex Sans", False, False), "Plex-Medium": ("IBM Plex Sans Medm", False, False),
          "Plex-SemiBold": ("IBM Plex Sans", True, False), "Plex-Bold": ("IBM Plex Sans", True, False),
          "Plex-Italic": ("IBM Plex Sans", False, True)}


def cargar(nombre, archivo):
    spec = importlib.util.spec_from_file_location(nombre, caminho(f"scripts/{archivo}"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def hexc(c):
    return c.hexval()[2:].upper() if hasattr(c, "hexval") else str(c).lstrip("#").upper()


# ------------------------------------------------------------------ registro de la maqueta
class Lienzo:
    """Sustituye al canvas de ReportLab y anota los elementos de cada página."""

    def __init__(self, tr=lambda s: s):
        self.tr = tr
        self.paginas, self.el = [], []
        self.fill, self.stroke, self.lw, self.font = "000000", "000000", 1, ("Plex", 10)

    def setFillColor(self, c): self.fill = hexc(c)
    def setStrokeColor(self, c): self.stroke = hexc(c)
    def setLineWidth(self, w): self.lw = w
    def setFont(self, f, s): self.font = (f, s)
    def setTitle(self, *a): pass
    setSubject = setAuthor = setKeywords = setTitle

    def rect(self, x, y, w, h, stroke=0, fill=1):
        self.el.append(dict(t="rect", x=x, y=y, w=w, h=h, fill=self.fill, r=0))

    def roundRect(self, x, y, w, h, r, stroke=0, fill=1):
        self.el.append(dict(t="rect", x=x, y=y, w=w, h=h, fill=self.fill, r=r))

    def line(self, x1, y1, x2, y2):
        self.el.append(dict(t="line", x1=x1, y1=y1, x2=x2, y2=y2, color=self.stroke, lw=self.lw))

    def drawString(self, x, y, s, align="l"):
        self.el.append(dict(t="str", x=x, y=y, s=self.tr(s), font=self.font, color=self.fill, align=align))

    def drawRightString(self, x, y, s):
        self.drawString(x, y, s, align="r")

    def drawImage(self, img, x, y, w, h, **k):
        self.el.append(dict(t="img", path=str(img.fileName), x=x, y=y, w=w, h=h))

    def showPage(self):
        self.paginas.append(self.el)
        self.el = []

    def save(self): pass


def parrafo(p):
    s = p.style
    return dict(text=p.text, estilo=s.name, font=s.fontName, size=s.fontSize, leading=s.leading,
                after=s.spaceAfter, before=s.spaceBefore, color=hexc(s.textColor), indent=s.leftIndent,
                bullet=getattr(p, "bulletText", None))


def instrumentar(mod, tr):
    """Sustituye las primitivas de dibujo del módulo 18 por versiones que registran."""
    from reportlab.platypus import Paragraph as ParReal, Table, TableStyle

    class ParRec(ParReal):
        def __init__(self, text, style, bulletText=None, **k):
            super().__init__(tr(text), style, bulletText=bulletText, **k)
            self.orig = text
            self.text = tr(text)
            self.bulletText = bulletText

        def drawOn(self, canvas, x, y, _sW=0):
            canvas.el.append(dict(t="pars", x=x, ytop=y + self.height, w=self.width, h=self.height, pars=[parrafo(self)]))

    mod.Paragraph = ParRec

    def texto(c, x, y_top, w, h, bloques):
        # Como en el PDF: el Frame descarta los bloques que ya no caben. Se mide con el texto original en español,
        # para que la versión traducida contenga exactamente los mismos bloques.
        import io
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.platypus import Frame
        resto = [ParReal(b.orig, b.style, bulletText=b.bulletText) for b in bloques]
        Frame(x, y_top - h, w, h, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0).addFromList(
            resto, Canvas(io.BytesIO()))
        dibujados = bloques[:len(bloques) - len(resto)]
        if resto and tr.lang == "es":
            print(f"  aviso: {len(resto)} bloque(s) no caben en su marco y tampoco figuran en el PDF: "
                  + " | ".join(r.text[:60] for r in resto))
        c.el.append(dict(t="pars", x=x, ytop=y_top, w=w, h=h, pars=[parrafo(b) for b in dibujados]))

    def figura(c, nombre, x, y_top, w, h, leyenda=None, alinear="top"):
        from reportlab.lib.utils import ImageReader
        ruta = (FIG_PT if tr.lang == "pt" else mod.FIG) / nombre
        iw, ih = ImageReader(str(mod.FIG / nombre)).getSize()   # la maqueta se basa en la figura original
        esc = min(w / iw, h / ih)
        dw, dh = iw * esc, ih * esc
        y = y_top - dh if alinear == "top" else y_top - h + (h - dh) / 2
        c.el.append(dict(t="img", path=str(ruta), x=x, y=y, w=dw, h=dh))
        if leyenda:
            p = ParRec(leyenda, mod.EST["leyenda"])
            _, ph = p.wrap(w, 60)
            c.el.append(dict(t="pars", x=x, ytop=y - 3, w=w, h=ph + 4, pars=[parrafo(p)]))
        return y

    def dato(c, x, y, valor, etiqueta, color=mod.ACENTO, w=110):
        c.el.append(dict(t="str", x=x, y=y, s=tr(valor), font=("Plex-SemiBold", 20), color=hexc(color), align="l"))
        p = ParRec(etiqueta, mod.ParagraphStyle("d", fontName="Plex", fontSize=7.4, leading=9.2, textColor=mod.SUAVE))
        _, ph = p.wrap(w, 40)
        c.el.append(dict(t="pars", x=x, ytop=y - 3, w=w, h=ph + 4, pars=[parrafo(p)]))

    def tabla(c, filas, x, y_top, anchos, cab_color=mod.ACENTO, extra=None):
        data = [[ParRec(str(v), mod.EST["tabla_b" if i == 0 else "tabla"]) for v in f] for i, f in enumerate(filas)]
        filas = [[p.text for p in f] for f in data]   # ya traducidas por ParRec
        t = Table(data, colWidths=anchos)
        t.setStyle(TableStyle([("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
                               ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
        _, th = t.wrap(sum(anchos), 400)
        # TEXTCOLOR de `extra` no afecta a celdas con Paragraph: en el PDF el texto conserva el color del estilo
        colores = {}
        c.el.append(dict(t="tabla", x=x, ytop=y_top, anchos=anchos, filas=filas, alturas=list(t._rowHeights),
                         cab=hexc(cab_color), colores={f"{k[0]},{k[1]}": v for k, v in colores.items()}))
        return y_top - th

    mod.texto, mod.figura, mod.dato, mod.tabla = texto, figura, dato, tabla


def registrar(lang):
    from traducao_pt import Tradutor
    tr = Tradutor(lang)
    mod = cargar("informe18", "18_informe_pdf.py")
    mod.fuentes()
    mod.estilos()
    d = mod.datos()
    instrumentar(mod, tr)
    c = Lienzo(tr)
    for f in (mod.portada, mod.p01, mod.p02, mod.p03, mod.p04, mod.p05, mod.p06, mod.p07, mod.p08, mod.p09, mod.p10,
              mod.p11, mod.p12):
        f(c, d)
    return mod, c.paginas, tr


# ------------------------------------------------------------------ PowerPoint
H_PT = 595.2755905511812
W_PT = 841.8897637795277


def E(pt):
    return Emu(int(round(pt * 12700)))


def runs_markup(texto):
    """'<b>x</b> y <font color='#c00'>z</font>' -> [(texto, bold, italic, color)]."""
    out, bold, ital, color = [], False, False, None
    for tok in re.split(r"(<[^>]+>)", texto):
        if not tok:
            continue
        if tok.startswith("<"):
            t = tok.lower()
            if t == "<b>":
                bold = True
            elif t == "</b>":
                bold = False
            elif t == "<i>":
                ital = True
            elif t == "</i>":
                ital = False
            elif t.startswith("<font"):
                m = re.search(r"color=['\"]#?([0-9a-fA-F]{6})", tok)
                color = m.group(1).upper() if m else color
            elif t == "</font>":
                color = None
            elif t in ("<br/>", "<br>"):
                out.append(("\n", bold, ital, color))
            continue
        tok = tok.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
        out.append((tok, bold, ital, color))
    return out


def fuente_run(run, font_rl, size, color, bold=False, ital=False):
    nombre, b0, i0 = FUENTE.get(font_rl, ("IBM Plex Sans", False, False))
    if bold and font_rl in ("Plex", "Plex-Medium"):
        # <b> en el PDF = Plex-SemiBold; como negrita de la familia, el énfasis se conserva aunque falte la fuente
        nombre, b0 = "IBM Plex Sans", True
    f = run.font
    f.name = nombre
    f.size = Pt(size)
    f.bold = b0
    f.italic = i0 or ital
    f.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", nombre)


def caja(slide, x, ytop, w, h, nombre=None, wrap=True):
    tb = slide.shapes.add_textbox(E(x), E(H_PT - ytop), E(w), E(max(h, 4)))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    if nombre:
        tb.name = nombre
    return tb, tf


def vineta(par, sangria):
    ppr = par._p.get_or_add_pPr()
    ppr.set("marL", str(int(sangria * 12700)))
    ppr.set("indent", str(int(-sangria * 12700)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in ppr.findall(qn(tag)):
            ppr.remove(e)
    bu = etree.SubElement(ppr, qn("a:buChar"))
    bu.set("char", "•")


def sin_vineta(par):
    ppr = par._p.get_or_add_pPr()
    etree.SubElement(ppr, qn("a:buNone"))


def poner_parrafos(tf, pars):
    primero = True
    for p in pars:
        par = tf.paragraphs[0] if primero else tf.add_paragraph()
        primero = False
        par.alignment = PP_ALIGN.LEFT
        par.line_spacing = Pt(p["leading"])
        par.space_after = Pt(p["after"] or 0)
        par.space_before = Pt(p["before"] or 0)
        if p["bullet"]:
            vineta(par, max(p["indent"], 10))
        else:
            sin_vineta(par)
        for txt, b, i, col in runs_markup(p["text"]):
            r = par.add_run()
            r.text = txt
            fuente_run(r, p["font"], p["size"], col or p["color"], bold=b, ital=i)


def borde(celda, lado, color, ancho):
    tcpr = celda._tc.get_or_add_tcPr()
    for e in tcpr.findall(qn(f"a:{lado}")):
        tcpr.remove(e)
    ln = etree.SubElement(tcpr, qn(f"a:{lado}"))
    if color is None:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    else:
        ln.set("w", str(int(ancho * 12700)))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
    # orden requerido por el esquema: lnL, lnR, lnT, lnB antes del relleno
    orden = ["lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D", "noFill", "solidFill", "gradFill",
             "blipFill", "pattFill", "grpFill", "headers", "extLst"]
    hijos = sorted(list(tcpr), key=lambda e: orden.index(etree.QName(e).localname) if etree.QName(e).localname in orden else 99)
    for e in hijos:
        tcpr.remove(e)
        tcpr.append(e)


def poner_tabla(slide, el, mod):
    filas, anchos = el["filas"], el["anchos"]
    gs = slide.shapes.add_table(len(filas), len(anchos), E(el["x"]), E(H_PT - el["ytop"]), E(sum(anchos)),
                                E(sum(el["alturas"])))
    gs.name = "Tabla"
    tbl = gs.table
    tbl._tbl.tblPr.set("firstRow", "0")
    tbl._tbl.tblPr.set("bandRow", "0")
    est = tbl._tbl.tblPr.find(qn("a:tableStyleId"))
    if est is None:
        est = etree.SubElement(tbl._tbl.tblPr, qn("a:tableStyleId"))
    est.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid
    for j, w in enumerate(anchos):
        tbl.columns[j].width = E(w)
    for i, h in enumerate(el["alturas"]):
        tbl.rows[i].height = E(h)
    for i, fila in enumerate(filas):
        st = mod.EST["tabla_b" if i == 0 else "tabla"]
        for j, v in enumerate(fila):
            cel = tbl.cell(i, j)
            cel.fill.background()
            cel.margin_left = cel.margin_right = E(3)
            cel.margin_top = cel.margin_bottom = E(2.2)
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cel.text_frame
            tf.word_wrap = True
            par = tf.paragraphs[0]
            par.line_spacing = Pt(st.leading)
            for txt, b, it, col in runs_markup(v):
                r = par.add_run()
                r.text = txt
                fuente_run(r, st.fontName, st.fontSize, el["colores"].get(f"{j},{i}") or col or hexc(st.textColor), bold=b, ital=it)
            borde(cel, "lnL", None, 0)
            borde(cel, "lnR", None, 0)
            borde(cel, "lnT", None, 0)
            borde(cel, "lnB", el["cab"] if i == 0 else "D5D9DD", 0.8 if i == 0 else 0.25)


def construir(lang):
    mod, paginas, tr = registrar(lang)
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W_PT), E(H_PT)
    blank = prs.slide_layouts[6]
    for n, els in enumerate(paginas):
        s = prs.slides.add_slide(blank)
        for el in els:
            if el["t"] == "rect":
                forma = MSO_SHAPE.ROUNDED_RECTANGLE if el["r"] else MSO_SHAPE.RECTANGLE
                sh = s.shapes.add_shape(forma, E(el["x"]), E(H_PT - el["y"] - el["h"]), E(el["w"]), E(el["h"]))
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor.from_string(el["fill"])
                sh.line.fill.background()
                sh.shadow.inherit = False
                if el["r"]:
                    sh.adjustments[0] = el["r"] / min(el["w"], el["h"])
                sh.name = "Fondo"
            elif el["t"] == "line":
                ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(el["x1"]), E(H_PT - el["y1"]), E(el["x2"]), E(H_PT - el["y2"]))
                ln.line.color.rgb = RGBColor.from_string(el["color"])
                ln.line.width = Pt(el["lw"])
                ln._element.spPr.append(ln._element.spPr.makeelement(qn("a:effectLst"), {}))   # sin sombra del tema
            elif el["t"] == "img":
                pic = s.shapes.add_picture(el["path"], E(el["x"]), E(H_PT - el["y"] - el["h"]), E(el["w"]), E(el["h"]))
                pic.name = Path(el["path"]).stem
            elif el["t"] == "str":
                fname, size = el["font"]
                ancho = stringWidth(el["s"], fname, size) * 1.12 + 4
                x = el["x"] - ancho if el["align"] == "r" else el["x"]
                tb, tf = caja(s, x, el["y"] + size * 1.0, ancho, size * 1.35, wrap=False)
                par = tf.paragraphs[0]
                par.alignment = PP_ALIGN.RIGHT if el["align"] == "r" else PP_ALIGN.LEFT
                r = par.add_run()
                r.text = el["s"]
                fuente_run(r, fname, size, el["color"])
            elif el["t"] == "pars":
                tb, tf = caja(s, el["x"], el["ytop"], el["w"], el["h"])
                poner_parrafos(tf, el["pars"])
            elif el["t"] == "tabla":
                poner_tabla(s, el, mod)
    cp = prs.core_properties
    cp.title = tr("Desigualdad territorial en la oferta de autobuses urbanos de Sevilla")
    cp.subject = tr("Densidad de población, accesibilidad peatonal y frecuencia programada del servicio TUSSAM")
    cp.author = "Proyecto sevilla_oferta_bus (autoría pendiente)"
    prs.save(SALIDA[lang])
    print(f"  {SALIDA[lang].relative_to(RAIZ)}: {len(paginas)} diapositivas")
    return tr


def figuras_pt():
    """Regenera las figuras con textos en portugués (sin tocar produtos/figuras/)."""
    import matplotlib.text as mtext
    from traducao_pt import Tradutor
    tr = Tradutor("pt")
    orig = mtext.Text.set_text

    def set_text(self, s):
        return orig(self, tr.fig(s) if isinstance(s, str) else s)
    mtext.Text.set_text = set_text
    try:
        v = cargar("visual", "visual.py")
        sys.modules["visual"] = v
        v.FIGURAS = FIG_PT
        f17 = cargar("figuras17", "17_figuras.py")
        f17.v = v
        f17.main()
    finally:
        mtext.Text.set_text = orig
    return tr


def coletar():
    """Vuelca a JSON todas las cadenas del informe y de las figuras (para mantener la traducción)."""
    import matplotlib.text as mtext
    from traducao_pt import Tradutor
    tr = Tradutor("es", coletar=True)
    mod = cargar("informe18", "18_informe_pdf.py")
    mod.fuentes()
    mod.estilos()
    d = mod.datos()
    instrumentar(mod, tr)
    c = Lienzo(tr)
    for f in (mod.portada, mod.p01, mod.p02, mod.p03, mod.p04, mod.p05, mod.p06, mod.p07, mod.p08, mod.p09, mod.p10,
              mod.p11, mod.p12):
        f(c, d)
    figs = []
    orig = mtext.Text.set_text

    def set_text(self, s):
        if isinstance(s, str) and s.strip():
            figs.append(s)
        return orig(self, s)
    mtext.Text.set_text = set_text
    v = cargar("visual", "visual.py")
    sys.modules["visual"] = v
    v.FIGURAS = FIG_PT   # se sobrescriben después con la versión traducida
    f17 = cargar("figuras17", "17_figuras.py")
    f17.v = v
    f17.main()
    mtext.Text.set_text = orig
    out = caminho("data/processed/textos_informe_es.json")
    out.write_text(json.dumps({"informe": sorted(set(tr.vistos)), "figuras": sorted(set(figs))}, ensure_ascii=False, indent=1),
                   encoding="utf-8")
    print(f"  {out.relative_to(RAIZ)}: {len(set(tr.vistos))} textos del informe, {len(set(figs))} de figuras")


def main():
    if "--coletar" in sys.argv:
        coletar()
        return
    construir("es")
    trf = figuras_pt()
    tr = construir("pt")
    for nombre, falt in (("informe", tr.faltantes), ("figuras", trf.faltantes)):
        if falt:
            print(f"  ATENCIÓN: {len(falt)} textos sin traducción ({nombre})")
            for t in sorted(falt)[:60]:
                print("   -", repr(t[:120]))


if __name__ == "__main__":
    main()
